#!/usr/bin/env python3
"""
Convert Job_Application_Tracker.xlsx into seed.json for the app's Import button.

Usage:
    python scripts/seed_from_xlsx.py [path/to/Job_Application_Tracker.xlsx]

Defaults to the tracker sitting in the JOB folder on this machine.
Reads the 'Pipeline' sheet: headers on row 3, data from row 4.
"""

import datetime
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is missing. Install it with:  pip install openpyxl")

DEFAULT_XLSX = os.path.join(
    os.path.expanduser("~"), "OneDrive", "Desktop", "JOB", "Job_Application_Tracker.xlsx"
)
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "seed.json")

VALID_STATUS = {
    "Not applied", "Applied", "Pending", "Heard back",
    "Interview", "Offer", "Rejected", "Withdrawn",
}
VALID_SOURCE = {
    "Seek", "LinkedIn", "Grad program", "Company site",
    "Referral", "Recruiter", "Other",
}
VALID_PRIORITY = {"High", "Medium", "Low"}


def iso(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    return None


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    if not os.path.exists(path):
        sys.exit(f"Could not find the workbook at:\n  {path}\n\nPass the path as an argument.")

    wb = openpyxl.load_workbook(path)
    if "Pipeline" not in wb.sheetnames:
        sys.exit(f"No 'Pipeline' sheet. Found: {', '.join(wb.sheetnames)}")
    ws = wb["Pipeline"]

    rows = []
    for r in range(4, ws.max_row + 1):
        company = ws.cell(r, 1).value
        if not company:
            continue

        link_cell = ws.cell(r, 4)
        link = link_cell.hyperlink.target if link_cell.hyperlink else ""

        status = ws.cell(r, 7).value or "Not applied"
        source = ws.cell(r, 5).value or "Other"
        priority = ws.cell(r, 6).value or "Medium"

        rows.append({
            "company": str(company).strip(),
            "role": str(ws.cell(r, 2).value or "").strip(),
            "location": str(ws.cell(r, 3).value or "").strip(),
            "link": link,
            "source": source if source in VALID_SOURCE else "Other",
            "priority": priority if priority in VALID_PRIORITY else "Medium",
            "status": status if status in VALID_STATUS else "Not applied",
            "date_found": iso(ws.cell(r, 8).value),
            "date_applied": iso(ws.cell(r, 9).value),
            "deadline": iso(ws.cell(r, 10).value),
            "last_update": iso(ws.cell(r, 11).value),
            "next_action": str(ws.cell(r, 12).value or "").strip(),
            # The spreadsheet has no due-date column, so infer one: an action
            # on a row that already has a deadline inherits it, otherwise the
            # app treats it as undated and it shows under "worth a look".
            "next_action_due": iso(ws.cell(r, 10).value) if ws.cell(r, 12).value else None,
            "notes": str(ws.cell(r, 13).value or "").strip(),
        })

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)

    counts = {}
    for row in rows:
        counts[row["status"]] = counts.get(row["status"], 0) + 1

    print(f"Wrote {len(rows)} roles to {OUT}\n")
    for status, n in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"  {n:>3}  {status}")
    print("\nNow open the app and use Import to load that file.")


if __name__ == "__main__":
    main()
